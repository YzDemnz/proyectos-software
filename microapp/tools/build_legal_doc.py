from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\MicroApp")
OUT_DIR = ROOT / "legal"
WEB_LEGAL_DIR = ROOT / "website" / "legal"
OUTPUT = OUT_DIR / "TERMINOS_Y_CONDICIONES_MICROAPP_BORRADOR.docx"
WEB_OUTPUT = WEB_LEGAL_DIR / OUTPUT.name

NAVY = "0B1B3A"
BLUE = "1464FF"
LIGHT_BLUE = "EAF1FF"
INK = "202938"
MUTED = "5D687A"
LIGHT_GRAY = "F3F5F8"
MID_GRAY = "D9E0EA"
GOLD_FILL = "FFF6DD"
GOLD_BORDER = "D8A72B"
RED_FILL = "FFF0F0"
RED = "A12626"
GREEN_FILL = "ECF9F1"
GREEN = "137A42"

# Resolved preset: contract_negotiation_brief.
# Letter portrait, 1-inch margins, Arial 11 pt, 1.25 line spacing,
# Heading 1 16 pt / Heading 2 13 pt / Heading 3 12 pt,
# table width 9360 DXA, table indent 120 DXA and explicit cell margins.


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bottom_border(paragraph, color: str = MID_GRAY, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run(run, 9, MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run(run, size: float = 11, color: str = INK, bold=False, italic=False) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(size)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def ensure_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element

    def create_abstract(num_id: int, fmt: str, text: str, left: int, hanging: int) -> int:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(num_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), str(left))
        indent.set(qn("w:hanging"), str(hanging))
        p_pr.append(tabs)
        p_pr.append(indent)
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        r_pr.append(fonts)
        level.extend([start, num_fmt, lvl_text, p_pr, r_pr])
        abstract.append(level)
        numbering.append(abstract)
        return num_id

    def create_num(num_id: int, abstract_id: int) -> int:
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId"))
    ]
    next_abstract = max(abstract_ids, default=0) + 1
    next_num = max(num_ids, default=0) + 1
    create_abstract(next_abstract, "bullet", "•", 540, 270)
    bullet_id = create_num(next_num, next_abstract)
    create_abstract(next_abstract + 1, "decimal", "%1.", 540, 270)
    decimal_id = create_num(next_num + 1, next_abstract + 1)
    return bullet_id, decimal_id


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run(first, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run(rest)
    else:
        run = p.add_run(text)
        set_run(run)
    return p


def add_bullet(doc: Document, text: str, bullet_id: int):
    p = add_body(doc, text)
    p.paragraph_format.space_after = Pt(4)
    apply_num(p, bullet_id)
    return p


def add_numbered(doc: Document, text: str, decimal_id: int):
    p = add_body(doc, text)
    p.paragraph_format.space_after = Pt(4)
    apply_num(p, decimal_id)
    return p


def add_callout(doc: Document, title: str, text: str, fill: str, border: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "8")
        edge.set(qn("w:color"), border)
        borders.append(edge)
    tc_pr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run(r, 10.5, border, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(text)
    set_run(r2, 10, INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run(r, 9.5, "FFFFFF", bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[idx], LIGHT_GRAY)
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(value)
            set_run(r, 9.2)
    for row in table.rows:
        row._tr.get_or_add_trPr()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 14, 8),
        ("Heading 2", 13, BLUE, 11, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def build_document() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    WEB_LEGAL_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    configure_styles(doc)
    bullet_id, decimal_id = ensure_numbering(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hr = hp.add_run("MICROAPP  |  TÉRMINOS Y CONDICIONES - BORRADOR")
    set_run(hr, 8.5, MUTED, bold=True)
    add_bottom_border(hp, MID_GRAY, 6)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)

    # Editorial cover.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MICROAPP")
    set_run(r, 12, BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Términos y condiciones de uso y contratación")
    set_run(r, 28, NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Borrador contractual robusto para revisión legal")
    set_run(r, 14, MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("Preparado para operación inicial en Torreón, Coahuila de Zaragoza, México")
    set_run(r, 11, NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Fecha de investigación normativa: 6 de junio de 2026")
    set_run(r, 10, MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Versión 0.9 - NO PUBLICAR SIN REVISIÓN DE ABOGADO")
    set_run(r, 10, RED, bold=True)

    doc.add_page_break()

    add_callout(
        doc,
        "AVISO DE REVISIÓN OBLIGATORIA",
        "Este documento es un borrador de trabajo elaborado para estructurar la relación contractual de MicroApp. "
        "No sustituye la opinión de un abogado mexicano ni debe publicarse hasta completar los datos del proveedor, "
        "validar el modelo de cobro, elaborar el aviso de privacidad y confirmar que las funciones reales de la app "
        "coinciden con lo prometido. Ninguna cláusula puede excluir derechos irrenunciables, responsabilidad por dolo, "
        "obligaciones de protección de datos, sanciones administrativas o responsabilidad penal.",
        RED_FILL,
        RED,
    )

    add_heading(doc, "Datos críticos que deben completarse", 1)
    add_table(
        doc,
        ["Campo", "Dato requerido antes de publicar"],
        [
            ["Proveedor", "[NOMBRE COMPLETO O RAZÓN SOCIAL]"],
            ["RFC", "[RFC DEL PROVEEDOR]"],
            ["Domicilio", "[DOMICILIO COMPLETO EN TORREÓN, COAHUILA]"],
            ["Soporte", "[CORREO Y TELÉFONO DE SOPORTE]"],
            ["Privacidad", "[CORREO DEL RESPONSABLE DE DATOS PERSONALES]"],
            ["Reclamaciones", "[MEDIO ELECTRÓNICO Y DOMICILIO PARA RECLAMACIONES]"],
            ["Dominio", "[DOMINIO DEFINITIVO DE MICROAPP]"],
            ["Cobro", "[PROCESADOR, PERIODICIDAD, IMPUESTOS Y POLÍTICA DE REEMBOLSO]"],
        ],
        [2700, 6660],
    )

    add_heading(doc, "Criterio de protección contractual", 1)
    add_body(
        doc,
        "El texto busca asignar responsabilidades de forma clara, limitar daños indirectos y evitar que MicroApp sea "
        "tratada como contador, banco, autoridad fiscal, aseguradora o garante de los negocios de sus usuarios. La "
        "protección más sólida no proviene de negar toda responsabilidad, sino de describir con precisión el servicio, "
        "sus dependencias, sus límites y los remedios disponibles, respetando la legislación de orden público.",
    )
    add_body(
        doc,
        "El Código Civil Federal establece que la voluntad privada no puede modificar normas de interés público y que "
        "la responsabilidad derivada de dolo no puede renunciarse. La Ley Federal de Protección al Consumidor protege "
        "derechos irrenunciables y permite tener por no puestas cláusulas abusivas. Por ello, toda limitación de este "
        "documento contiene la frase \"en la máxima medida permitida por la ley\" y exclusiones obligatorias.",
    )

    add_heading(doc, "Índice de revisión", 1)
    toc_items = [
        "1. Identidad del proveedor y definiciones",
        "2. Aceptación electrónica, capacidad y cuenta",
        "3. Alcance funcional de MicroApp",
        "4. Planes, precios, cobro recurrente y cancelación",
        "5. Datos del usuario, privacidad y seguridad",
        "6. Inventario, ventas, tickets y estimaciones",
        "7. Deudas, mensajería, automatizaciones y terceros",
        "8. Mercado Pago Point y servicios externos",
        "9. Propiedad intelectual y disponibilidad",
        "10. Garantías, límites de responsabilidad e indemnidad",
        "11. Suspensión, terminación, disputas y ley aplicable",
        "Anexos A-F: planes, uso aceptable, datos, incidentes, implementación y fuentes",
    ]
    for item in toc_items:
        add_bullet(doc, item, bullet_id)

    doc.add_page_break()

    add_heading(doc, "TÉRMINOS Y CONDICIONES DE MICROAPP", 1)
    add_body(
        doc,
        "Última actualización del borrador: 6 de junio de 2026. Estos Términos regulan el acceso, descarga, registro, "
        "contratación y uso de la aplicación móvil, aplicación web, sitio público, integraciones, soporte y servicios "
        "relacionados identificados comercialmente como MicroApp.",
    )

    add_heading(doc, "1. Identidad del proveedor", 1)
    add_body(
        doc,
        "MicroApp es operada por [NOMBRE COMPLETO O RAZÓN SOCIAL], con RFC [RFC], domicilio en [DOMICILIO COMPLETO], "
        "Torreón, Coahuila de Zaragoza, México, correo de atención [CORREO] y teléfono [TELÉFONO], en adelante "
        "“MicroApp”, “el Proveedor”, “nosotros” o “nuestro”. La información de identificación deberá mostrarse de "
        "manera clara antes de contratar cualquier plan de pago.",
    )

    add_heading(doc, "2. Definiciones", 1)
    definitions = [
        ("Aplicación", "MicroApp para dispositivos móviles, web y cualquier módulo asociado."),
        ("Cuenta", "Registro individual o empresarial desde el cual se accede al Servicio."),
        ("Negocio", "Actividad comercial o profesional cuyos datos administra el Usuario."),
        ("Usuario", "Persona física que accede a la Aplicación por cuenta propia o en representación de un Negocio."),
        ("Titular", "Persona que creó la cuenta principal y controla la configuración del Negocio."),
        ("Datos del Usuario", "Información, archivos, fotografías, inventarios, ventas, clientes, deudas y demás contenido cargado o generado por el Usuario."),
        ("Servicio", "Funciones de organización comercial disponibles según el plan contratado."),
        ("Plan Gratis", "Acceso sin costo con las funciones y límites que se muestren en la Aplicación."),
        ("Plan Negocio", "Suscripción mensual que habilita funciones avanzadas conforme a la pantalla de contratación."),
        ("Proveedor Externo", "Tercero independiente, como una tienda de aplicaciones, servicio de nube, Meta/WhatsApp o Mercado Pago."),
        ("Movimiento", "Registro trazable que explica una entrada, salida, reserva, ajuste, devolución o merma de inventario."),
        ("Ticket", "Comprobante operativo generado por MicroApp que no constituye CFDI salvo integración expresa con un proveedor autorizado."),
    ]
    add_table(doc, ["Término", "Significado"], [[a, b] for a, b in definitions], [2500, 6860])

    add_heading(doc, "3. Aceptación electrónica", 1)
    add_body(
        doc,
        "Al seleccionar una casilla de aceptación, pulsar “Crear cuenta”, “Empezar”, “Contratar”, “Continuar”, "
        "“Aceptar” o una acción equivalente, el Usuario manifiesta su consentimiento a estos Términos y reconoce que "
        "tuvo oportunidad razonable de leerlos, descargarlos y conservarlos. La aceptación se registrará con fecha, "
        "hora, versión del documento, identificador de cuenta y demás evidencia técnica proporcional.",
    )
    add_body(
        doc,
        "Los mensajes de datos y consentimientos electrónicos producirán efectos en los términos del Código de Comercio "
        "y del Código Civil Federal cuando la información sea atribuible, íntegra y accesible para consulta posterior. "
        "MicroApp conservará evidencia de la versión aceptada, sin perjuicio de los derechos del Usuario para solicitarla.",
    )
    add_body(
        doc,
        "Si el Usuario no acepta estos Términos, no deberá crear una Cuenta ni usar el Servicio. La navegación meramente "
        "informativa por el sitio público se sujetará a las disposiciones que por su naturaleza resulten aplicables.",
    )

    add_heading(doc, "4. Capacidad y representación", 1)
    add_body(
        doc,
        "El Usuario declara que tiene al menos dieciocho años, capacidad legal suficiente y, cuando actúe por un "
        "Negocio o persona moral, facultades para obligarla. MicroApp podrá solicitar información razonable para "
        "verificar identidad, representación o titularidad, especialmente ante solicitudes de recuperación, exportación "
        "masiva, cambio de Titular o ejercicio de derechos sobre datos personales.",
    )
    add_body(
        doc,
        "La Cuenta no está diseñada para ser contratada directamente por menores de edad. Si MicroApp detecta una Cuenta "
        "creada sin capacidad suficiente, podrá restringirla mientras verifica la autorización o representación legal, "
        "sin apropiarse de los Datos del Usuario.",
    )

    add_heading(doc, "5. Registro y seguridad de la Cuenta", 1)
    for text in [
        "El Usuario proporcionará información veraz, actual y completa, y mantendrá actualizados sus datos de contacto.",
        "Las credenciales son personales. El Titular es responsable de usar contraseñas robustas, proteger su dispositivo y cerrar sesiones en equipos ajenos.",
        "El Usuario notificará de inmediato cualquier acceso no autorizado, pérdida del dispositivo, sospecha de fraude o cambio de personal con acceso.",
        "MicroApp podrá aplicar autenticación adicional, bloqueo temporal, límites de intentos, verificación de dispositivo y otras medidas proporcionales.",
        "El Titular será responsable de las acciones realizadas por empleados o invitados autorizados, sin perjuicio de la responsabilidad individual de quien actúe ilícitamente.",
    ]:
        add_bullet(doc, text, bullet_id)
    add_body(
        doc,
        "MicroApp nunca solicitará por chat público la contraseña completa, códigos de autenticación o claves privadas. "
        "El Usuario no deberá compartir tokens, credenciales de Mercado Pago ni secretos de integraciones con empleados "
        "que no requieran acceso.",
    )

    add_heading(doc, "6. Naturaleza y alcance del Servicio", 1)
    add_body(
        doc,
        "MicroApp es una herramienta tecnológica de apoyo administrativo. Permite registrar y consultar información "
        "relacionada con productos, categorías, existencias, ventas, pagos, clientes, deudas, abonos, tickets, "
        "movimientos, reportes, empleados, proveedores, citas, catálogos e integraciones, según el plan y la etapa de "
        "desarrollo disponible.",
    )
    add_body(
        doc,
        "Las funciones visibles en materiales promocionales deberán identificarse como disponibles, beta, próximas o "
        "dependientes de terceros. Una descripción comercial no garantiza que todas las funciones estén habilitadas en "
        "todas las plataformas, países, dispositivos, giros o planes.",
    )
    add_callout(
        doc,
        "REGLA DE PRODUCTO",
        "MicroApp no deberá anunciar como operativa una integración que aún requiera configuración, certificación, "
        "autorización del tercero o desarrollo adicional. Las promesas públicas deben coincidir con la función real.",
        GOLD_FILL,
        GOLD_BORDER,
    )

    add_heading(doc, "7. Registros operativos y folios", 1)
    add_body(
        doc,
        "Cada venta confirmada podrá generar un folio, detalle de productos, registro de pago, movimiento de inventario "
        "y ticket. El Usuario reconoce que estos registros dependen de la información capturada, de la correcta "
        "configuración de productos y de la sincronización disponible. El folio organiza operaciones internas y no "
        "convierte por sí mismo el Ticket en comprobante fiscal.",
    )
    add_body(
        doc,
        "Los productos asociados a operaciones históricas deberán archivarse en lugar de eliminarse cuando su supresión "
        "afecte la trazabilidad. MicroApp podrá impedir o advertir sobre eliminaciones que rompan referencias, sin asumir "
        "responsabilidad por ajustes manuales autorizados por el Usuario.",
    )

    add_heading(doc, "8. Planes y funciones", 1)
    add_body(
        doc,
        "Las funciones, límites y precios vigentes se mostrarán antes de la contratación. El Plan Gratis ofrece una "
        "muestra funcional y podrá incluir inventario, ventas, clientes, deudas, tickets, exportaciones y reportes "
        "básicos. El Plan Negocio podrá habilitar catálogo, empleados, permisos, multi-dispositivo, compras, proveedores, "
        "agenda, automatizaciones, integraciones y reportes avanzados.",
    )
    add_body(
        doc,
        "MicroApp podrá mejorar, reorganizar o sustituir funciones, pero no reducirá de manera sustancial una prestación "
        "pagada durante el periodo ya cubierto sin ofrecer una alternativa razonable, ajuste proporcional, migración o "
        "cancelación conforme a la ley. Los límites se aplicarán de forma clara y no retroactiva sobre datos ya capturados.",
    )

    add_heading(doc, "9. Precio, impuestos y comprobación del pago", 1)
    add_body(
        doc,
        "Los precios se expresarán en pesos mexicanos e indicarán si incluyen impuestos. El Usuario autoriza únicamente "
        "los cargos mostrados en la pantalla final de contratación. El procesador de pagos podrá solicitar información "
        "adicional y aplicar sus propios controles antifraude.",
    )
    add_body(
        doc,
        "La activación del plan ocurre cuando MicroApp o el Proveedor Externo recibe confirmación válida del pago. Una "
        "captura de pantalla, mensaje no verificado o cargo pendiente no obliga a activar el Servicio hasta que exista "
        "confirmación suficiente. Si un cargo es revertido, desconocido o fraudulento, MicroApp podrá suspender las "
        "funciones pagadas mientras se aclara, preservando los datos en la medida razonable.",
    )

    add_heading(doc, "10. Renovación automática y cancelación", 1)
    add_body(
        doc,
        "Toda renovación automática requerirá consentimiento expreso e informado. Antes de contratar se mostrará de "
        "forma clara la periodicidad, monto, fecha estimada de cobro, impuestos aplicables, mecanismo de cancelación y "
        "alcance del plan. El consentimiento para cobros recurrentes se guardará separado de la aceptación general "
        "cuando sea técnicamente posible.",
    )
    add_body(
        doc,
        "MicroApp notificará cada renovación automática con al menos cinco días naturales de anticipación. La "
        "notificación indicará el importe previsto, la fecha de cargo, el plan y un acceso directo al mecanismo de "
        "cancelación. El Usuario podrá cancelar de manera inmediata y sin penalización para evitar cargos futuros.",
    )
    add_body(
        doc,
        "La cancelación no borra automáticamente la Cuenta ni los Datos del Usuario. Salvo que la ley, la tienda de "
        "aplicaciones o una política más favorable disponga otra cosa, el acceso pagado continuará hasta el final del "
        "periodo vigente. Si el Usuario solicita cierre inmediato y eliminación, se le advertirá sobre las consecuencias "
        "y se respetarán obligaciones de conservación legal.",
    )
    add_body(
        doc,
        "Cuando la suscripción sea facturada por Google Play u otra tienda, el Usuario podrá administrar la renovación "
        "en dicha plataforma. MicroApp deberá ofrecer instrucciones claras y no obstaculizar la cancelación.",
    )

    add_heading(doc, "11. Pruebas gratuitas, promociones y recompensas", 1)
    add_body(
        doc,
        "Las pruebas gratuitas indicarán duración, funciones incluidas y qué sucede al terminar. Una prueba no se "
        "convertirá en plan pagado sin consentimiento expreso para el cobro. Las promociones por uso continuo, como una "
        "prueba de treinta días después de quince días consecutivos de actividad, podrán depender de criterios objetivos "
        "de elegibilidad, prevención de abuso y disponibilidad.",
    )
    add_body(
        doc,
        "MicroApp podrá cancelar promociones obtenidas mediante cuentas duplicadas, automatización artificial, "
        "suplantación o manipulación, sin afectar derechos adquiridos de buena fe. Las promociones no son canjeables por "
        "efectivo y no garantizan elegibilidad futura.",
    )

    add_heading(doc, "12. Cambios de precio", 1)
    add_body(
        doc,
        "Los cambios de precio se aplicarán a periodos futuros y se comunicarán antes de la renovación con tiempo "
        "razonable y por un medio que permita conservar la notificación. Si el Usuario no acepta el nuevo precio, podrá "
        "cancelar antes del siguiente cargo. No se cobrarán importes distintos sin el consentimiento requerido.",
    )

    add_heading(doc, "13. Titularidad de los Datos del Usuario", 1)
    add_body(
        doc,
        "El Usuario conserva la titularidad y los derechos que le correspondan sobre los Datos del Usuario. Al usar el "
        "Servicio concede a MicroApp una licencia limitada, no exclusiva, revocable al terminar la relación y únicamente "
        "por el tiempo necesario para alojar, procesar, respaldar, transmitir y mostrar dichos datos con el fin de prestar, "
        "proteger y mejorar el Servicio.",
    )
    add_body(
        doc,
        "MicroApp no adquiere propiedad sobre inventarios, listas de clientes, historiales, fotografías, marcas o "
        "documentos del Usuario. La licencia no autoriza vender bases de clientes ni utilizar secretos comerciales para "
        "competir con el Negocio. El uso estadístico deberá realizarse de forma agregada o disociada cuando corresponda.",
    )

    add_heading(doc, "14. Aviso de privacidad y datos personales", 1)
    add_body(
        doc,
        "El tratamiento de datos personales se regirá por un Aviso de Privacidad separado, accesible antes de recabar "
        "datos. Ese aviso deberá identificar al responsable, domicilio, finalidades, categorías de datos, opciones para "
        "limitar uso, transferencias, medios para ejercer derechos y procedimiento para comunicar cambios.",
    )
    add_body(
        doc,
        "Estos Términos no sustituyen el Aviso de Privacidad. Cuando exista conflicto, se aplicará la disposición que "
        "otorgue mayor claridad y protección a la persona titular, sin modificar obligaciones legales. MicroApp tratará "
        "datos bajo principios de licitud, finalidad, lealtad, consentimiento, calidad, proporcionalidad, información y "
        "responsabilidad.",
    )
    add_body(
        doc,
        "El Usuario podrá solicitar acceso, rectificación, cancelación u oposición, así como revocar consentimiento o "
        "limitar el uso, mediante [CORREO DE PRIVACIDAD]. La solicitud deberá atenderse conforme a los plazos, excepciones "
        "y procedimientos de la Ley Federal de Protección de Datos Personales en Posesión de los Particulares.",
    )

    add_heading(doc, "15. Datos personales de clientes, empleados y terceros", 1)
    add_body(
        doc,
        "El Usuario puede cargar nombres, teléfonos, notas, adeudos, citas, fotografías u otros datos de terceros. El "
        "Usuario declara que cuenta con base jurídica, aviso, consentimiento o autorización suficiente para recabarlos, "
        "usarlos, conservarlos y comunicar mensajes. No deberá capturar datos sensibles que no sean estrictamente "
        "necesarios ni información obtenida mediante engaño, compra ilícita o extracción no autorizada.",
    )
    add_body(
        doc,
        "Dependiendo del flujo, el Usuario puede actuar como responsable de los datos de sus clientes y MicroApp como "
        "encargado tecnológico. Esta asignación deberá precisarse en un acuerdo de tratamiento de datos cuando el plan "
        "Negocio incluya servicios de nube, empleados, automatización o procesamiento sistemático por cuenta del Usuario.",
    )

    add_heading(doc, "16. Seguridad de la información", 1)
    add_body(
        doc,
        "MicroApp implementará medidas administrativas, técnicas y físicas razonables conforme a la naturaleza de los "
        "datos, riesgos, estado de la técnica y capacidad operativa. Podrán incluir cifrado en tránsito, cifrado o "
        "protección de secretos, control de acceso, registros de actividad, respaldos, límites de solicitudes, gestión "
        "de vulnerabilidades y separación de entornos.",
    )
    add_body(
        doc,
        "Ningún sistema es absolutamente invulnerable. MicroApp no garantiza seguridad perfecta, pero esta aclaración no "
        "la libera de aplicar medidas exigibles, corregir fallas bajo su control ni notificar vulneraciones que afecten "
        "significativamente derechos patrimoniales o morales cuando la ley lo requiera.",
    )
    add_body(
        doc,
        "El Usuario colaborará razonablemente ante incidentes: cambiará credenciales, revocará sesiones, preservará "
        "evidencia y evitará publicar datos sensibles o detalles que aumenten el riesgo antes de contener la vulneración.",
    )

    add_heading(doc, "17. Respaldo, sincronización, exportación y conservación", 1)
    add_body(
        doc,
        "MicroApp podrá ofrecer almacenamiento local, sincronización y respaldo en nube. El Usuario entiende que la "
        "sincronización puede retrasarse por conectividad, batería, permisos, capacidad del dispositivo, servicios "
        "externos o mantenimiento. Debe revisar el estado de sincronización antes de cambiar de equipo, borrar la app o "
        "restablecer el dispositivo.",
    )
    add_body(
        doc,
        "Cuando esté disponible, el Usuario podrá exportar información en formatos como CSV, Excel, Word o PDF. La "
        "compatibilidad depende de la estructura del archivo y de las aplicaciones externas. MicroApp podrá rechazar "
        "archivos corruptos, ejecutables, macros, contenido malicioso o formatos que no pueda interpretar de manera segura.",
    )
    add_body(
        doc,
        "Tras terminar la Cuenta, MicroApp conservará los datos por el plazo informado para permitir exportación y cumplir "
        "obligaciones legales, atender disputas, prevenir fraude y mantener respaldos rotativos. Después podrá eliminarlos "
        "o disociarlos de forma segura. Los plazos definitivos deberán declararse en el Aviso de Privacidad.",
    )

    add_heading(doc, "18. Importación y clasificación asistida", 1)
    add_body(
        doc,
        "La importación de inventario y la clasificación automática por categorías son herramientas de asistencia. Pueden "
        "interpretar de forma incorrecta nombres ambiguos, unidades, precios, impuestos o categorías. El Usuario debe "
        "revisar la vista previa antes de confirmar. MicroApp no es responsable por errores que el Usuario confirme sin "
        "revisión, salvo cuando deriven directamente de una falla imputable que la ley obligue a reparar.",
    )
    add_body(
        doc,
        "El Usuario no cargará documentos con datos innecesarios, contraseñas, información bancaria completa, macros, "
        "código malicioso ni contenido cuya utilización infrinja derechos de terceros.",
    )

    add_heading(doc, "19. Inventario, existencias y movimientos", 1)
    add_body(
        doc,
        "Las existencias mostradas son el resultado de saldos iniciales, ventas, compras, reservas, devoluciones, mermas y "
        "ajustes registrados. MicroApp no realiza conteos físicos ni puede detectar por sí sola robo, daño, caducidad, "
        "errores de captura o mercancía no registrada. El Usuario realizará conciliaciones periódicas.",
    )
    add_body(
        doc,
        "Toda modificación automatizada de stock deberá generar un Movimiento trazable. Los ajustes manuales identificarán "
        "al usuario, fecha, cantidad y motivo cuando la función esté disponible. MicroApp podrá impedir stock negativo o "
        "mostrar advertencias, pero la configuración final depende del giro y del Usuario.",
    )

    add_heading(doc, "20. Ventas, pagos, deudas y abonos", 1)
    add_body(
        doc,
        "Una venta podrá pasar por estados como borrador, pendiente, pagada, fiada, parcial, cancelada, reembolsada o "
        "devuelta. El Usuario es responsable de seleccionar el estado correcto y verificar que la entrega, cobro, "
        "devolución y movimiento de inventario coincidan con la operación real.",
    )
    add_body(
        doc,
        "El registro de una deuda o abono es una herramienta administrativa y no prueba por sí solo la existencia, "
        "exigibilidad o vencimiento legal de una obligación frente a terceros. El Usuario debe conservar contratos, "
        "notas, comprobantes y consentimientos que resulten necesarios.",
    )

    add_heading(doc, "21. Tickets, impuestos y contabilidad", 1)
    add_body(
        doc,
        "Los tickets generados por MicroApp son resúmenes operativos. No son CFDI, facturas fiscales, recibos de nómina, "
        "contabilidad oficial ni declaraciones ante el SAT, salvo que una integración futura indique expresamente que "
        "interviene un proveedor autorizado y cumpla los requisitos aplicables.",
    )
    add_body(
        doc,
        "Los reportes de ventas, costos, utilidad, caja e impuestos son estimaciones basadas en la información capturada. "
        "No sustituyen asesoría contable, fiscal, legal o financiera. El Usuario debe validar resultados con sus registros "
        "y asesores antes de tomar decisiones, presentar declaraciones o responder a autoridades.",
    )

    add_heading(doc, "22. Mensajería y cobranza", 1)
    add_body(
        doc,
        "Las funciones de WhatsApp, SMS, correo u otros canales podrán generar mensajes prellenados, recordatorios o "
        "automatizaciones. El Usuario es el remitente y responsable del contenido, destinatario, frecuencia, horario, "
        "consentimiento, identificación comercial, mecanismos de oposición y cumplimiento de las reglas del canal.",
    )
    for text in [
        "No se enviarán amenazas, hostigamiento, discriminación, lenguaje engañoso ni afirmaciones falsas sobre consecuencias legales.",
        "No se revelarán deudas a familiares, empleadores, grupos o terceros no autorizados.",
        "Los recordatorios automáticos deberán poder pausarse y respetar solicitudes de no recibir mensajes.",
        "El Usuario revisará montos, nombre del cliente, fecha y contexto antes de enviar mensajes sensibles.",
        "MicroApp podrá limitar envíos ante señales de abuso, quejas, rebotes, bloqueo del canal o riesgo legal.",
    ]:
        add_bullet(doc, text, bullet_id)

    add_heading(doc, "23. Automatizaciones y funciones asistidas", 1)
    add_body(
        doc,
        "Las reglas automáticas y asistentes internos pueden sugerir categorías, respuestas, recordatorios, resúmenes o "
        "acciones. Son herramientas de apoyo y pueden producir resultados incompletos o equivocados. El Usuario mantendrá "
        "revisión humana para decisiones que afecten pagos, cobranza, derechos, empleo, inventario relevante o "
        "comunicaciones externas.",
    )
    add_body(
        doc,
        "MicroApp no autoriza automatizaciones destinadas a suplantar personas, engañar clientes, recolectar datos sin "
        "permiso, evadir bloqueos, enviar spam, discriminar o tomar decisiones legales autónomas. Podrá desactivar una "
        "automatización riesgosa mientras solicita corrección.",
    )

    add_heading(doc, "24. Integración con Mercado Pago Point", 1)
    add_body(
        doc,
        "Cuando se habilite, el flujo integrado podrá crear una venta pendiente, enviar una orden a una terminal, esperar "
        "confirmación del procesador y completar la venta únicamente después de recibir un estado válido. El inventario "
        "podrá reservarse temporalmente y descontarse de forma definitiva al confirmarse el pago.",
    )
    add_body(
        doc,
        "Mercado Pago es un Proveedor Externo independiente. MicroApp no custodia fondos, no autoriza tarjetas, no decide "
        "contracargos, no garantiza disponibilidad de terminales y no sustituye los términos, comisiones, verificaciones "
        "o políticas de Mercado Pago. El Usuario deberá vincular su propia cuenta y proteger sus credenciales.",
    )
    add_body(
        doc,
        "Los webhooks, consultas y referencias podrán retrasarse, duplicarse o llegar fuera de orden. MicroApp implementará "
        "controles de idempotencia y conciliación razonables, pero el Usuario deberá revisar operaciones pendientes antes "
        "de entregar bienes cuando el pago no esté confirmado. Un estado visual no sustituye la conciliación del "
        "procesador ante disputas.",
    )
    add_body(
        doc,
        "Los reembolsos y devoluciones requieren dos decisiones distintas: devolver dinero y, cuando el producto haya "
        "regresado físicamente en condiciones aceptables, reingresarlo a inventario. MicroApp no aumentará stock "
        "automáticamente por un reembolso si no existe confirmación de devolución del producto.",
    )

    add_heading(doc, "25. Servicios y contenido de terceros", 1)
    add_body(
        doc,
        "MicroApp puede depender de tiendas de aplicaciones, sistemas operativos, nube, mapas, cámara, WhatsApp/Meta, "
        "Mercado Pago, proveedores de correo y otros terceros. Esos servicios se rigen por sus propios términos y pueden "
        "cambiar, suspender o limitar APIs. MicroApp no responde por actos exclusivamente imputables al tercero, pero "
        "realizará esfuerzos razonables para comunicar afectaciones relevantes y ofrecer alternativas cuando sea viable.",
    )
    add_body(
        doc,
        "Los enlaces externos no implican patrocinio ni control. El Usuario evaluará las políticas del tercero antes de "
        "conectar una cuenta o transferir datos. MicroApp solicitará solo los permisos necesarios y deberá explicar su "
        "finalidad en la interfaz.",
    )

    add_heading(doc, "26. Propiedad intelectual", 1)
    add_body(
        doc,
        "La Aplicación, código, diseño, marcas, bases estructurales, textos, interfaces, documentación y elementos "
        "propios de MicroApp están protegidos por la legislación aplicable. La suscripción concede una licencia limitada, "
        "revocable, no exclusiva, no transferible y sin derecho de sublicencia para usar el Servicio conforme a estos "
        "Términos.",
    )
    add_body(
        doc,
        "El Usuario no podrá descompilar salvo en casos expresamente permitidos por ley, extraer bases, eludir límites, "
        "copiar interfaces de forma sustancial, revender acceso sin autorización, eliminar avisos, usar marcas de forma "
        "confusa ni realizar pruebas de seguridad destructivas. La investigación responsable podrá reportarse a "
        "[CORREO DE SEGURIDAD].",
    )

    add_heading(doc, "27. Contenido, logos y fotografías del Usuario", 1)
    add_body(
        doc,
        "El Usuario declara tener derechos o autorización para cargar logos, fotografías, nombres comerciales, catálogos "
        "y descripciones. MicroApp podrá comprimir o adaptar técnicamente los archivos sin alterar su sentido. No está "
        "obligada a revisar de forma general su legalidad, pero podrá retirar contenido ante una reclamación fundada, "
        "orden competente o infracción manifiesta.",
    )

    add_heading(doc, "28. Disponibilidad, mantenimiento y versiones beta", 1)
    add_body(
        doc,
        "MicroApp procurará una disponibilidad razonable, sin prometer operación ininterrumpida. Puede realizar "
        "mantenimiento, actualizaciones, migraciones o medidas urgentes de seguridad. Cuando sea posible, comunicará "
        "interrupciones programadas que afecten funciones esenciales.",
    )
    add_body(
        doc,
        "Las funciones beta, piloto o de acceso anticipado pueden cambiar, contener errores o retirarse. Se identificarán "
        "como tales y no deberán usarse como único mecanismo para una operación crítica sin respaldo alternativo.",
    )

    add_heading(doc, "29. Soporte", 1)
    add_body(
        doc,
        "El soporte se prestará por los canales y horarios publicados. Los tiempos de respuesta son objetivos y no "
        "garantías, salvo acuerdo de nivel de servicio separado. El Usuario proporcionará información suficiente y no "
        "enviará contraseñas, datos bancarios completos ni información personal que no sea necesaria para resolver el caso.",
    )

    add_heading(doc, "30. Obligaciones y usos prohibidos", 1)
    prohibited = [
        "Usar el Servicio para fraude, suplantación, acceso ilícito, lavado, evasión, amenazas, extorsión o cualquier delito.",
        "Alterar pagos, folios, tickets o registros con el propósito de engañar a clientes, autoridades, socios o empleados.",
        "Acceder o intentar acceder a cuentas, sistemas, datos, APIs o credenciales sin autorización.",
        "Introducir malware, macros dañinas, código de explotación o archivos destinados a afectar el Servicio.",
        "Realizar envíos masivos no solicitados, hostigamiento de cobranza o tratamiento ilícito de datos personales.",
        "Revender, alquilar o compartir una Cuenta fuera del alcance permitido por el plan.",
        "Interferir con límites, medición de uso, seguridad, facturación, controles de plan o disponibilidad.",
        "Usar el Servicio para tomar decisiones discriminatorias o prohibidas sin revisión humana y base legal.",
    ]
    for item in prohibited:
        add_bullet(doc, item, bullet_id)
    add_body(
        doc,
        "MicroApp podrá conservar evidencia y colaborar con autoridades cuando exista obligación legal, orden válida o "
        "necesidad de proteger derechos, aplicando controles de proporcionalidad y privacidad. Estos Términos no crean "
        "inmunidad frente a responsabilidad civil, administrativa o penal.",
    )

    add_heading(doc, "31. Declaraciones y garantías limitadas", 1)
    add_body(
        doc,
        "MicroApp declara que prestará el Servicio con cuidado razonable, conforme a la descripción aplicable y sin "
        "ocultar intencionalmente defectos conocidos relevantes. Si una función pagada no se presta de manera sustancial, "
        "el Usuario deberá notificarlo para permitir diagnóstico, corrección, repetición del servicio, ajuste o remedio "
        "legal correspondiente.",
    )
    add_body(
        doc,
        "En la máxima medida permitida por la ley, no se garantiza que el Servicio satisfaga todas las necesidades "
        "particulares, elimine todo error humano, impida toda pérdida, aumente ventas, cobre deudas, produzca resultados "
        "fiscales correctos o sea compatible con todo dispositivo y servicio externo.",
    )

    add_heading(doc, "32. Limitación de responsabilidad", 1)
    add_body(
        doc,
        "En la máxima medida permitida por la ley, MicroApp no será responsable por daños indirectos, incidentales, "
        "especiales o consecuenciales, lucro cesante remoto, pérdida de oportunidades, daño reputacional o pérdida de "
        "datos que no sea consecuencia inmediata y directa de un incumplimiento imputable a MicroApp.",
    )
    add_body(
        doc,
        "La responsabilidad contractual acumulada de MicroApp por reclamaciones relacionadas con un mismo hecho se "
        "limitará al mayor entre: (a) las cantidades efectivamente pagadas por el Usuario durante los tres meses "
        "inmediatamente anteriores al hecho; o (b) $1,000.00 MXN cuando el Usuario utilice exclusivamente un plan gratuito. "
        "Este límite solo operará cuando sea legalmente válido y podrá ajustarse por un abogado para evitar carácter abusivo.",
    )
    add_body(
        doc,
        "La limitación anterior no aplica a: dolo; responsabilidad que por ley no pueda excluirse; afectaciones a vida o "
        "integridad física imputables; incumplimiento deliberado de obligaciones de confidencialidad; tratamiento ilícito "
        "de datos personales atribuible a MicroApp; infracción intencional de propiedad intelectual de terceros; ni "
        "remedios irrenunciables de la persona consumidora.",
    )
    add_body(
        doc,
        "MicroApp no responderá por decisiones del Usuario basadas en datos incompletos, falta de conteo físico, "
        "credenciales compartidas, mensajes enviados sin revisión, archivos mal estructurados, dispositivos inseguros, "
        "servicios de terceros o incumplimientos legales propios del Negocio, salvo en la proporción en que MicroApp haya "
        "contribuido directamente al daño.",
    )

    add_heading(doc, "33. Indemnidad limitada", 1)
    add_body(
        doc,
        "El Usuario mantendrá indemne a MicroApp frente a reclamaciones de terceros que deriven directamente de contenido "
        "ilícito cargado por el Usuario, falta de autorización sobre datos de sus clientes, cobranza abusiva, uso criminal, "
        "infracción de propiedad intelectual o incumplimiento de obligaciones propias del Negocio. Esta obligación se "
        "limita a montos razonables y comprobables, permite al Usuario participar en la defensa y no cubre la conducta "
        "culposa o dolosa atribuible a MicroApp.",
    )
    add_body(
        doc,
        "MicroApp notificará oportunamente la reclamación y no celebrará acuerdos que impongan admisión de culpa u "
        "obligaciones no monetarias al Usuario sin su consentimiento razonable.",
    )

    add_heading(doc, "34. Caso fortuito y fuerza mayor", 1)
    add_body(
        doc,
        "Ninguna parte responderá por retrasos causados por eventos imprevisibles o inevitables fuera de su control "
        "razonable, como desastres, fallas generales de telecomunicaciones, actos de autoridad, conflictos laborales, "
        "ciberataques generalizados o interrupciones graves de proveedores. La parte afectada mitigará daños, informará "
        "cuando sea posible y reanudará el cumplimiento. No se invocará fuerza mayor cuando la parte causó o contribuyó "
        "al evento, aceptó expresamente el riesgo o la ley imponga responsabilidad.",
    )

    add_heading(doc, "35. Suspensión y medidas de protección", 1)
    add_body(
        doc,
        "MicroApp podrá suspender temporalmente funciones cuando exista riesgo de fraude, seguridad, impago confirmado, "
        "uso ilícito, afectación a terceros o incumplimiento material. Cuando el riesgo lo permita, notificará la causa y "
        "dará oportunidad razonable de corregir. En emergencias podrá actuar de inmediato y explicar después.",
    )
    add_body(
        doc,
        "La suspensión será proporcional: se procurará bloquear la función o integración afectada antes que toda la Cuenta. "
        "MicroApp no usará la suspensión para apropiarse de datos, impedir exportaciones legítimas o presionar pagos "
        "controvertidos de mala fe.",
    )

    add_heading(doc, "36. Terminación y salida de datos", 1)
    add_body(
        doc,
        "El Usuario puede dejar de usar el Servicio, cancelar la suscripción o solicitar cierre de Cuenta. Antes de cerrar, "
        "deberá exportar la información que necesite. MicroApp ofrecerá un periodo de recuperación razonable de [30 DÍAS] "
        "después del cierre, salvo que el Usuario solicite eliminación inmediata, exista obligación de conservar o una "
        "restricción de seguridad.",
    )
    add_body(
        doc,
        "Al terminar, cesan las licencias de uso, pero sobreviven las cláusulas que por su naturaleza deban continuar: "
        "propiedad, pagos devengados, confidencialidad, responsabilidad, indemnidad, conservación legal, solución de "
        "controversias y protección de datos.",
    )

    add_heading(doc, "37. Reembolsos y remedios", 1)
    add_body(
        doc,
        "Salvo derecho legal, política de la tienda o promesa comercial más favorable, la cancelación evita cobros futuros "
        "y no genera reembolso proporcional del periodo ya iniciado. Sí procederá revisión cuando exista cargo no "
        "autorizado, cobro duplicado, servicio no prestado de forma sustancial, incumplimiento de garantía o supuesto "
        "previsto por la legislación de consumo.",
    )
    add_body(
        doc,
        "El Usuario presentará la reclamación con folio, fecha, importe y evidencia disponible. MicroApp no condicionará "
        "derechos legales a formalidades excesivas y proporcionará respuesta clara por un medio durable.",
    )

    add_heading(doc, "38. Modificaciones a los Términos", 1)
    add_body(
        doc,
        "MicroApp podrá actualizar estos Términos por cambios legales, seguridad, nuevas funciones o modelo comercial. "
        "Las modificaciones materiales se comunicarán antes de entrar en vigor. No se aplicarán retroactivamente para "
        "eliminar derechos ya causados ni se presumirá aceptación de cambios sustanciales por silencio cuando la ley "
        "requiera consentimiento.",
    )
    add_body(
        doc,
        "Cuando el cambio afecte precio, cobro recurrente, tratamiento de datos no compatible o una prestación esencial, "
        "se solicitará la aceptación correspondiente o se permitirá cancelar sin penalización.",
    )

    add_heading(doc, "39. Comunicaciones y notificaciones", 1)
    add_body(
        doc,
        "MicroApp podrá enviar comunicaciones operativas a los datos registrados: seguridad, recuperación, cambios "
        "contractuales, renovaciones, cargos, soporte y continuidad. Las comunicaciones comerciales requerirán las "
        "preferencias o consentimientos aplicables y ofrecerán medios para dejar de recibirlas.",
    )
    add_body(
        doc,
        "El Usuario mantendrá actualizado su correo y teléfono. Una notificación no se considerará inválida solo porque "
        "el Usuario no revisó un canal actualizado y razonablemente utilizado, sin perjuicio de requisitos especiales "
        "para cobros, privacidad o actos que exijan confirmación.",
    )

    add_heading(doc, "40. Atención de quejas y solución de controversias", 1)
    add_body(
        doc,
        "Antes de iniciar una controversia, las partes procurarán resolverla mediante soporte en [CORREO], proporcionando "
        "una descripción concreta y una propuesta de solución. MicroApp acusará recibo y responderá dentro de un plazo "
        "razonable según la complejidad.",
    )
    add_body(
        doc,
        "La persona consumidora conserva el derecho de acudir a la Procuraduría Federal del Consumidor y, cuando resulte "
        "disponible y aplicable, a mecanismos electrónicos de conciliación. Ninguna negociación privada limita el acceso "
        "a autoridades competentes, medidas urgentes o derechos irrenunciables.",
    )

    add_heading(doc, "41. Ley aplicable y jurisdicción", 1)
    add_body(
        doc,
        "Estos Términos se regirán por las leyes federales mexicanas y, de manera supletoria cuando corresponda, por las "
        "leyes del Estado de Coahuila de Zaragoza. Para controversias que legalmente puedan someterse a jurisdicción "
        "convencional, las partes señalan los tribunales competentes de Torreón, Coahuila, renunciando a otro fuero que "
        "pudiera corresponder por domicilio presente o futuro.",
    )
    add_body(
        doc,
        "La elección anterior no priva a la persona consumidora de acudir a PROFECO, ejercer derechos en el lugar que "
        "determinen normas imperativas ni utilizar cualquier foro que legalmente no pueda renunciarse. Tampoco desplaza "
        "competencias administrativas o penales.",
    )

    add_heading(doc, "42. Cesión, separabilidad, no renuncia y acuerdo completo", 1)
    add_body(
        doc,
        "El Usuario no podrá ceder una Cuenta sin autorización cuando la cesión comprometa seguridad, licencias o datos. "
        "MicroApp podrá ceder el contrato por reorganización, venta de negocio o sucesión, siempre que el adquirente "
        "asuma obligaciones y se informe cuando la ley lo requiera.",
    )
    add_body(
        doc,
        "Si una cláusula es inválida, se interpretará o reducirá al máximo permitido y las demás continuarán vigentes, "
        "salvo que la disposición sea esencial y su eliminación cambie injustamente el equilibrio. La falta de ejercicio "
        "de un derecho no implica renuncia. Estos Términos, el Aviso de Privacidad, la pantalla de contratación y anexos "
        "aplicables integran el acuerdo.",
    )

    add_heading(doc, "43. Orden de prelación", 1)
    add_body(
        doc,
        "En caso de conflicto se aplicará, de mayor a menor prioridad: (1) normas imperativas; (2) condiciones específicas "
        "mostradas y aceptadas al contratar; (3) acuerdo de tratamiento de datos o nivel de servicio, si existe; (4) estos "
        "Términos; (5) documentación operativa. Ningún material publicitario reducirá una garantía expresamente ofrecida.",
    )

    add_heading(doc, "44. Entrada en vigor y evidencia de aceptación", 1)
    add_body(
        doc,
        "La versión definitiva entrará en vigor en la fecha publicada. MicroApp conservará un historial de versiones y "
        "permitirá consultar la versión aplicable al momento de contratación. La evidencia podrá incluir aceptación, "
        "fecha, IP truncada o protegida, dispositivo, plan, consentimiento de renovación y copia de la pantalla final, "
        "aplicando minimización y plazos de conservación.",
    )

    doc.add_page_break()
    add_heading(doc, "ANEXO A. Matriz de planes y compromisos comerciales", 1)
    add_table(
        doc,
        ["Materia", "Gratis", "Negocio mensual"],
        [
            ["Precio", "$0 MXN", "$199 MXN, salvo promoción informada"],
            ["Inventario y ventas", "Uso operativo sin límite artificial anunciado", "Incluido"],
            ["Clientes, deudas y abonos", "Funciones base", "Funciones ampliadas y automatización"],
            ["Reportes", "Resumen y reportes básicos", "Reportes avanzados y corte ampliado"],
            ["Tickets", "Ticket operativo", "Personalización ampliada"],
            ["Empleados y permisos", "No incluido", "Incluido según disponibilidad"],
            ["Multi-dispositivo", "Limitado", "Incluido según disponibilidad"],
            ["Catálogo y mensajes", "Vista o función limitada", "Incluido según disponibilidad"],
            ["Mercado Pago Point", "No incluido", "Integración sujeta a configuración y tercero"],
            ["Cancelación", "No aplica cobro", "Inmediata y sin penalización para cargos futuros"],
        ],
        [2300, 3150, 3910],
    )
    add_body(
        doc,
        "La matriz debe actualizarse antes de publicar y cada vez que cambie la oferta. La leyenda “según disponibilidad” "
        "no debe utilizarse para cobrar por una función inexistente; sirve únicamente para funciones en despliegue "
        "gradual claramente identificadas antes de contratar.",
    )

    add_heading(doc, "ANEXO B. Política de uso aceptable", 1)
    acceptable_sections = [
        ("Seguridad", "No intentar acceso no autorizado, ingeniería social, extracción masiva, explotación destructiva ni evasión de controles."),
        ("Datos", "No cargar datos obtenidos ilícitamente, datos sensibles innecesarios, bases compradas sin autorización o secretos de terceros."),
        ("Mensajes", "No enviar spam, hostigamiento, amenazas, contenido discriminatorio ni cobranza a contactos no relacionados."),
        ("Pagos", "No simular aprobaciones, manipular webhooks, alterar referencias ni entregar tickets falsos."),
        ("Propiedad intelectual", "No copiar, revender, retirar marcas, usar logos ajenos sin permiso ni distribuir versiones modificadas."),
        ("Legalidad", "No usar MicroApp para fraude, evasión, comercio prohibido, lavado, suplantación o cualquier conducta delictiva."),
    ]
    add_table(doc, ["Área", "Regla"], [[a, b] for a, b in acceptable_sections], [2200, 7160])

    add_heading(doc, "ANEXO C. Bases mínimas de tratamiento de datos", 1)
    for text in [
        "Definir por escrito quién es responsable y quién es encargado en cada flujo.",
        "Publicar Aviso de Privacidad integral y simplificado antes de recabar datos.",
        "Documentar finalidades primarias y separar finalidades comerciales opcionales.",
        "Obtener consentimiento expreso cuando la ley lo exija, especialmente para datos sensibles o transferencias no exceptuadas.",
        "Permitir derechos de acceso, rectificación, cancelación y oposición con identidad verificada.",
        "Mantener inventario de datos, plazos de conservación, controles de acceso y registro de proveedores.",
        "Firmar acuerdos con nube, mensajería, analítica y soporte que procesen datos por cuenta de MicroApp.",
        "Preparar respuesta a incidentes y notificación cuando una vulneración afecte significativamente derechos.",
        "Evitar almacenar tokens de terceros en el celular cuando puedan protegerse en backend.",
        "Aplicar minimización: no pedir datos que la función no necesita.",
    ]:
        add_bullet(doc, text, bullet_id)

    add_heading(doc, "ANEXO D. Protocolo contractual de incidentes", 1)
    incident_steps = [
        "Detectar, registrar y clasificar el incidente sin alterar evidencia necesaria.",
        "Contener credenciales, sesiones, integraciones o rutas afectadas.",
        "Evaluar datos involucrados, titulares, alcance, probabilidad e impacto.",
        "Corregir la causa y validar que la medida no destruya información necesaria.",
        "Notificar internamente a responsable legal, seguridad, soporte y dirección.",
        "Informar a titulares y autoridades cuando corresponda, con lenguaje claro y medidas de protección.",
        "Documentar decisiones, tiempos, proveedores involucrados y acciones posteriores.",
        "Actualizar controles, contratos, capacitación y pruebas para evitar recurrencia.",
    ]
    for item in incident_steps:
        add_numbered(doc, item, decimal_id)

    add_heading(doc, "ANEXO E. Lista de implementación antes del lanzamiento", 1)
    implementation = [
        "Constituir o identificar correctamente al proveedor y completar RFC, domicilio y contacto.",
        "Registrar o evaluar la marca MicroApp y revisar disponibilidad de dominio y paquete de aplicación.",
        "Publicar Términos, Aviso de Privacidad y pantalla de consentimiento en versiones descargables.",
        "Separar aceptación de Términos, consentimiento de privacidad y autorización de cobro recurrente.",
        "Programar aviso de renovación con mínimo cinco días naturales.",
        "Programar cancelación inmediata y comprobante de cancelación.",
        "Crear historial de versiones y evidencia de aceptación.",
        "Definir retención, exportación, eliminación y recuperación de datos.",
        "Revisar contratos con nube, correo, mensajería, analítica y pagos.",
        "Configurar respaldo, cifrado de secretos, rotación, registros y respuesta a incidentes.",
        "Validar que tickets indiquen claramente que no son CFDI.",
        "Revisar textos publicitarios para evitar promesas absolutas de seguridad, ganancias o disponibilidad.",
        "Probar reembolsos, contracargos, pagos duplicados, webhooks repetidos y ventas pendientes.",
        "Habilitar canal de reclamaciones y procedimiento ARCO.",
        "Solicitar revisión final de abogado con experiencia en consumo, tecnología y datos personales en México.",
    ]
    for item in implementation:
        add_bullet(doc, f"☐ {item}", bullet_id)

    add_heading(doc, "ANEXO F. Matriz normativa consultada", 1)
    normative_rows = [
        ["Código Civil Federal", "Consentimiento, forma, obligaciones, daños, dolo, fuerza mayor y límites contractuales.", "Última reforma DOF 14-11-2025."],
        ["Código de Comercio", "Contratación mercantil y mensajes de datos.", "Texto vigente; montos actualizados DOF 18-02-2026."],
        ["Ley Federal de Protección al Consumidor", "Información, comercio electrónico, adhesión, garantías, cobro recurrente y cancelación.", "Última reforma DOF 12-12-2025."],
        ["LFPDPPP", "Tratamiento, aviso, seguridad, vulneraciones, derechos y transferencias.", "Nueva ley DOF 20-03-2025; reforma DOF 14-11-2025."],
        ["Código Penal Federal", "Fraude y acceso ilícito a sistemas; las obligaciones penales no se excluyen por contrato.", "Última reforma DOF 13-03-2026."],
        ["Código Civil de Coahuila", "Normas supletorias locales, renuncias, contratos y responsabilidad.", "Compilación del Congreso consultada; reforma indicada 27-06-2023."],
        ["Código Penal de Coahuila", "Fraude y conductas locales relevantes.", "Código publicado 27-10-2017 y reformas integradas en compilación consultada."],
    ]
    add_table(doc, ["Ordenamiento", "Relevancia", "Corte consultado"], normative_rows, [2450, 4310, 2600])

    add_heading(doc, "Fuentes oficiales", 2)
    sources = [
        ("Cámara de Diputados - Código Civil Federal", "https://www.diputados.gob.mx/LeyesBiblio/pdf/CCF.pdf"),
        ("Cámara de Diputados - Código de Comercio", "https://www.diputados.gob.mx/LeyesBiblio/pdf/CCom.pdf"),
        ("Cámara de Diputados - Ley Federal de Protección al Consumidor", "https://www.diputados.gob.mx/LeyesBiblio/pdf/LFPC.pdf"),
        ("Cámara de Diputados - Ley Federal de Protección de Datos Personales en Posesión de los Particulares", "https://www.diputados.gob.mx/LeyesBiblio/pdf/LFPDPPP.pdf"),
        ("Cámara de Diputados - Código Penal Federal", "https://www.diputados.gob.mx/LeyesBiblio/pdf/CPF.pdf"),
        ("Cámara de Diputados - índice de legislación federal vigente", "https://www.diputados.gob.mx/LeyesBiblio/index.htm"),
        ("Congreso de Coahuila - leyes estatales vigentes", "https://www.congresocoahuila.gob.mx/portal/leyes-estatales-vigentes/"),
        ("Congreso de Coahuila - Código Civil", "https://congresocoahuila.gob.mx/transparencia/03/Leyes_Coahuila/coa02.pdf"),
        ("Congreso de Coahuila - Código Penal 2017", "https://congresocoahuila.gob.mx/transparencia/03/Leyes_Coahuila/coa08_Nuevo_Codigo.pdf"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(f"{label}: ")
        set_run(r, 9, NAVY, bold=True)
        add_hyperlink(p, url, url)

    add_callout(
        doc,
        "NOTA SOBRE LEGISLACIÓN LOCAL",
        "El propio Congreso de Coahuila advierte que su compilación es informativa y que la publicación oficial en el "
        "Periódico Oficial determina la validez normativa. El abogado revisor deberá verificar reformas posteriores, "
        "criterios judiciales y entrada gradual del Código Nacional de Procedimientos Civiles y Familiares en el distrito "
        "judicial aplicable antes de emitir una versión definitiva.",
        GOLD_FILL,
        GOLD_BORDER,
    )

    add_heading(doc, "Cláusulas que requieren decisión del propietario", 1)
    decisions = [
        ["Identidad jurídica", "Operar como persona física con actividad empresarial o constituir persona moral."],
        ["Precio y cobro", "Confirmar $199 MXN, impuestos, fecha de cargo y procesador."],
        ["Renovación", "Definir canal de aviso, prueba de entrega y cancelación en un toque."],
        ["Retención", "Definir días de recuperación, respaldos y eliminación."],
        ["Límite de responsabilidad", "Validar monto, periodo y compatibilidad con LFPC."],
        ["Jurisdicción", "Confirmar domicilio y conveniencia de Torreón, con reserva de derechos del consumidor."],
        ["Privacidad", "Designar responsable, correo ARCO, proveedores y finalidades."],
        ["Mercado Pago", "Confirmar modalidad de integración, cuenta, términos y manejo de webhooks."],
        ["Mensajería", "Definir si será manual, semiautomática o automática y cómo se obtiene consentimiento."],
    ]
    add_table(doc, ["Tema", "Decisión pendiente"], decisions, [2700, 6660])

    add_heading(doc, "Constancia interna de revisión", 1)
    add_body(doc, "Propietario / representante: ______________________________________________")
    add_body(doc, "Abogado revisor: ________________________________________________________")
    add_body(doc, "Fecha de revisión: ______________________________________________________")
    add_body(doc, "Versión aprobada para publicación: _______________________________________")
    add_body(doc, "Observaciones: __________________________________________________________")
    add_body(doc, "________________________________________________________________________")

    core = doc.core_properties
    core.title = "Términos y condiciones de MicroApp - Borrador"
    core.subject = "Borrador contractual para revisión legal"
    core.author = "MicroApp"
    core.keywords = "MicroApp, términos, condiciones, México, Coahuila, SaaS"
    core.comments = "Documento de trabajo. No publicar sin revisión legal."

    doc.save(OUTPUT)
    WEB_OUTPUT.write_bytes(OUTPUT.read_bytes())
    print(str(OUTPUT))
    print(str(WEB_OUTPUT))


if __name__ == "__main__":
    build_document()
