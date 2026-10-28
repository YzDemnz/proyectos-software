import argparse
import hashlib
import json
import re
import shutil
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CP_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
DCTERMS_NS = "http://purl.org/dc/terms/"

NAMESPACES = {
    "w": W_NS,
    "cp": CP_NS,
    "dc": DC_NS,
    "dcterms": DCTERMS_NS,
}

LEGAL_FILENAME_WORDS = {
    "accion",
    "proforma",
    "aclaracion",
    "amparo",
    "directo",
    "indirecto",
    "contestacion",
    "convenio",
    "cumplimiento",
    "prevencion",
    "demanda",
    "custodia",
    "medida",
    "pension",
    "pesion",
    "denuncia",
    "deshaucio",
    "desahucio",
    "devolucion",
    "afore",
    "divorcio",
    "incausado",
    "violencia",
    "familiar",
    "escrito",
    "requejo",
    "lerdo",
    "fase",
    "alegatos",
    "manifiesta",
    "sobre",
    "peritaje",
    "rendido",
    "oficio",
    "imss",
    "pliego",
    "posiciones",
    "civil",
    "pruebas",
    "ordinario",
    "queja",
    "rebeldia",
    "recosideracion",
    "reconsideracion",
    "corrija",
    "sentencia",
    "interlocutriia",
    "interlocutoria",
    "solicita",
    "audio",
    "video",
    "escritura",
    "local",
    "junta",
    "tribunal",
}

INSTITUTION_WORDS = {
    "juzgado",
    "primero",
    "segundo",
    "tercero",
    "cuarto",
    "quinto",
    "sexto",
    "septimo",
    "octavo",
    "noveno",
    "decimo",
    "civil",
    "mercantil",
    "familiar",
    "penal",
    "laboral",
    "primera",
    "instancia",
    "distrito",
    "colegiado",
    "supremo",
    "tribunal",
    "sala",
    "unitario",
    "circuito",
    "torreon",
    "coahuila",
    "durango",
    "lerdo",
    "poder",
    "judicial",
    "federacion",
    "ministerio",
    "publico",
    "centro",
    "conciliacion",
}

LEGAL_HEADING_WORDS = INSTITUTION_WORDS | {
    "accion",
    "amparo",
    "antecedentes",
    "articulo",
    "articulos",
    "audiencia",
    "civil",
    "codigo",
    "conceptos",
    "contestacion",
    "contra",
    "demanda",
    "derecho",
    "directo",
    "escrito",
    "expediente",
    "federal",
    "fraccion",
    "fundamento",
    "garantias",
    "hechos",
    "incidente",
    "indirecto",
    "inicial",
    "juicio",
    "justicia",
    "laboral",
    "materia",
    "mercantil",
    "necesario",
    "ordinario",
    "peticion",
    "petitoria",
    "presente",
    "prestaciones",
    "primera",
    "primero",
    "proforma",
    "promocion",
    "protesto",
    "pruebas",
    "queja",
    "recurso",
    "sentencia",
    "solicito",
    "suma",
    "unico",
    "violacion",
    "convenio",
    "circuito",
    "considerando",
    "conforme",
    "concederse",
    "condena",
    "cumplir",
    "judicial",
    "octavo",
    "ofrecen",
    "opera",
    "primera",
    "quinto",
    "resolutivo",
    "termino",
    "trabajo",
}

SENSITIVE_PATTERNS = [
    ("email", re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)),
    ("credential", re.compile(r"\b[A-Za-zÁÉÍÓÚÑáéíóúñ]{5,}\d{2,}\b", re.IGNORECASE)),
    ("curp", re.compile(r"\b[A-Z][AEIOUX][A-Z]{2}\d{6}[HM][A-Z]{5}[A-Z0-9]\d\b", re.IGNORECASE)),
    ("rfc", re.compile(r"\b[A-ZÑ&]{3,4}\d{6}[A-Z0-9]{3}\b", re.IGNORECASE)),
    ("case_number", re.compile(r"\b\d{1,7}\s*/\s*\d{2,4}\b")),
    ("long_number", re.compile(r"\b\d{8,}\b")),
    ("phone", re.compile(r"(?<!\d)(?:\+?52[\s.-]?)?(?:\(?\d{2,3}\)?[\s.-]?)?\d{3,4}[\s.-]?\d{4}(?!\d)")),
    (
        "address",
        re.compile(
            r"\b(?:domicilio|calle|avenida|av\.|colonia|fraccionamiento|fracc\.|numero exterior|numero interior|c\.?\s*p\.?)\b[^.;\n]{0,140}",
            re.IGNORECASE,
        ),
    ),
    (
        "named_context",
        re.compile(
            r"\b(?:se llama|nombres?|identifica como|conocid[oa] como)\s+[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑa-záéíóúñ]+(?:\s+(?:y|e)\s+[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑa-záéíóúñ]+)?",
            re.IGNORECASE,
        ),
    ),
]

CONTEXT_NAME_PATTERN = re.compile(
    r"(?P<prefix>\b(?:C\.?|CC\.?|LIC\.?|LICENCIAD[OA]|SR\.?|SRA\.?|SEÑOR[AO]?|ACTOR[AO]?|DEMANDAD[OA]|QUEJOS[AO]|PROMOVENTE|APODERAD[OA]|DENUNCIANTE|IMPUTAD[OA])\s+)"
    r"(?P<name>[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑa-záéíóúñ]+(?:\s+[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑa-záéíóúñ]+){1,5})"
)

UPPERCASE_NAME_SEQUENCE = re.compile(r"\b[A-ZÁÉÍÓÚÑ]{3,}(?:\s+[A-ZÁÉÍÓÚÑ]{3,}){1,9}\b")


def normalize(value: str) -> str:
    return "".join(
        char for char in unicodedata.normalize("NFD", value)
        if unicodedata.category(char) != "Mn"
    ).lower()


def slugify(value: str) -> str:
    normalized = normalize(value)
    normalized = re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")
    return normalized or "escrito"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def infer_filename_terms(files: list[Path]) -> list[str]:
    terms: set[str] = set()
    for path in files:
        words = [
            normalize(word)
            for word in re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]{3,}", path.stem)
        ]
        possible_names = [word for word in words if word not in LEGAL_FILENAME_WORDS]
        for word in possible_names:
            if word not in INSTITUTION_WORDS:
                terms.add(word)
        for size in (2, 3):
            for index in range(0, max(0, len(possible_names) - size + 1)):
                chunk = possible_names[index:index + size]
                if all(word not in INSTITUTION_WORDS for word in chunk):
                    terms.add(" ".join(chunk))
    return sorted(terms, key=lambda item: (-len(item), item))


def make_name_patterns(filename_terms: list[str]) -> list[re.Pattern[str]]:
    accent_map = {
        "a": "[aáàäâã]",
        "e": "[eéèëê]",
        "i": "[iíìïî]",
        "o": "[oóòöôõ]",
        "u": "[uúùüû]",
        "n": "[nñ]",
    }

    def flexible(term: str) -> str:
        pieces = []
        for char in term:
            if char == " ":
                pieces.append(r"\s+")
            elif char in accent_map:
                pieces.append(accent_map[char])
            else:
                pieces.append(re.escape(char))
        return "".join(pieces)

    patterns = []
    for term in filename_terms:
        patterns.append(re.compile(rf"\b{flexible(term)}\b", re.IGNORECASE))
    return patterns


def mask_fragment(value: str) -> str:
    chars = []
    for char in value:
        if char.isspace():
            chars.append(char)
        elif char in ".,;:/\\-()[]{}'\"":
            chars.append(char)
        else:
            chars.append("X")
    return "".join(chars)


def is_likely_private_name(value: str) -> bool:
    words = [normalize(word) for word in re.findall(r"[A-ZÁÉÍÓÚÑ]{3,}", value)]
    if len(words) < 2:
        return False
    if any(word in LEGAL_HEADING_WORDS for word in words):
        return False
    if len(words) > 10:
        return False
    return True


def repair_mojibake(value: str) -> str:
    if "Ã" not in value and "â" not in value:
        return value
    try:
        candidate = value.encode("latin1").decode("utf-8")
    except UnicodeError:
        return value
    bad_original = value.count("Ã") + value.count("â")
    bad_candidate = candidate.count("Ã") + candidate.count("â")
    return candidate if bad_candidate < bad_original else value


def redact_text(value: str, name_patterns: list[re.Pattern[str]], placeholder: bool = False) -> tuple[str, Counter]:
    counts: Counter = Counter()

    def replace_match(label: str, replacement: str):
        def inner(match: re.Match[str]) -> str:
            counts[label] += 1
            return replacement if placeholder else mask_fragment(match.group(0))
        return inner

    text = value
    for label, pattern in SENSITIVE_PATTERNS:
        replacement = {
            "email": "[CORREO]",
            "curp": "[CURP]",
            "rfc": "[RFC]",
            "case_number": "[EXPEDIENTE]",
            "credential": "[CREDENCIAL]",
            "long_number": "[FOLIO]",
            "named_context": "[NOMBRE]",
            "phone": "[TELEFONO]",
            "address": "[DOMICILIO]",
        }[label]
        text = pattern.sub(replace_match(label, replacement), text)

    def replace_context_name(match: re.Match[str]) -> str:
        counts["context_name"] += 1
        if placeholder:
            return f"{match.group('prefix')}[NOMBRE]"
        return f"{match.group('prefix')}{mask_fragment(match.group('name'))}"

    text = CONTEXT_NAME_PATTERN.sub(replace_context_name, text)

    for pattern in name_patterns:
        text = pattern.sub(replace_match("filename_name", "[NOMBRE]"), text)

    def replace_uppercase_name(match: re.Match[str]) -> str:
        if not is_likely_private_name(match.group(0)):
            return match.group(0)
        counts["uppercase_name"] += 1
        return "[NOMBRE]" if placeholder else mask_fragment(match.group(0))

    text = UPPERCASE_NAME_SEQUENCE.sub(replace_uppercase_name, text)

    return text, counts


def redact_xml_part(data: bytes, name_patterns: list[re.Pattern[str]]) -> tuple[bytes, Counter]:
    parser = etree.XMLParser(remove_blank_text=False, recover=True)
    root = etree.fromstring(data, parser=parser)
    counts: Counter = Counter()

    for paragraph in root.xpath(".//w:p", namespaces=NAMESPACES):
        text_nodes = paragraph.xpath(".//w:t", namespaces=NAMESPACES)
        if not text_nodes:
            continue
        original_parts = [node.text or "" for node in text_nodes]
        combined = "".join(original_parts)
        redacted, current_counts = redact_text(combined, name_patterns, placeholder=False)
        if redacted == combined:
            continue
        counts.update(current_counts)
        cursor = 0
        for node, original in zip(text_nodes, original_parts):
            node.text = redacted[cursor:cursor + len(original)]
            cursor += len(original)

    for element in root.iter():
        for attr_name in list(element.attrib):
            if attr_name.endswith("rsid") or "rsid" in attr_name:
                del element.attrib[attr_name]

    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), counts


def scrub_core_properties(data: bytes) -> bytes:
    parser = etree.XMLParser(remove_blank_text=False, recover=True)
    root = etree.fromstring(data, parser=parser)
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    replacements = {
        f"{{{DC_NS}}}creator": "Judicial Managment",
        f"{{{CP_NS}}}lastModifiedBy": "Judicial Managment",
        f"{{{DCTERMS_NS}}}created": now,
        f"{{{DCTERMS_NS}}}modified": now,
    }
    for element in root.iter():
        if element.tag in replacements:
            element.text = replacements[element.tag]
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def sanitize_docx(source: Path, destination: Path, name_patterns: list[re.Pattern[str]]) -> Counter:
    counts: Counter = Counter()
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == "docProps/custom.xml":
                counts["custom_metadata_removed"] += 1
                continue

            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = scrub_core_properties(data)
            elif item.filename.startswith("word/") and item.filename.endswith(".xml"):
                data, part_counts = redact_xml_part(data, name_patterns)
                counts.update(part_counts)
            elif item.filename.endswith(".rels"):
                data, part_counts = redact_relationships(data)
                counts.update(part_counts)

            zout.writestr(item, data)
    return counts


def redact_relationships(data: bytes) -> tuple[bytes, Counter]:
    parser = etree.XMLParser(remove_blank_text=False, recover=True)
    root = etree.fromstring(data, parser=parser)
    counts: Counter = Counter()
    for element in root.iter():
        target = element.attrib.get("Target")
        if target and re.search(r"mailto:|@", target, re.IGNORECASE):
            element.attrib["Target"] = "#"
            counts["relationship_email"] += 1
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), counts


def extract_docx_text(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def infer_kind_and_matter(title: str, text: str) -> tuple[str, str, list[str]]:
    haystack = normalize(f"{title} {text[:2500]}")
    rules = [
        ("amparo", "amparo", "federal", ["amparo", "garantias", "acto reclamado"]),
        ("queja", "recurso de queja", "federal", ["queja", "recurso", "agravio"]),
        ("contestacion", "contestacion", "civil", ["contestacion", "excepciones", "defensas"]),
        ("demanda custodia", "demanda familiar", "familiar", ["custodia", "pension", "convivencia", "menores"]),
        ("divorcio", "demanda familiar", "familiar", ["divorcio", "violencia familiar", "medidas"]),
        ("convenio", "convenio", "familiar", ["convenio", "convivencia", "alimentos"]),
        ("denuncia", "denuncia", "penal", ["denuncia", "hechos", "ministerio publico"]),
        ("deshaucio", "demanda civil", "civil", ["desahucio", "arrendamiento", "desocupacion"]),
        ("desahucio", "demanda civil", "civil", ["desahucio", "arrendamiento", "desocupacion"]),
        ("proforma", "demanda mercantil", "mercantil", ["accion proforma", "escritura", "compraventa"]),
        ("afore", "promocion seguridad social", "laboral", ["afore", "imss", "seguridad social"]),
        ("alegatos", "alegatos", "civil", ["alegatos", "pruebas", "sentencia"]),
        ("peritaje", "manifestacion", "civil", ["peritaje", "dictamen", "objecion"]),
        ("pliego", "pliego de posiciones", "civil", ["confesional", "posiciones", "prueba"]),
        ("pruebas", "ofrecimiento de pruebas", "civil", ["pruebas", "testimonial", "documental"]),
        ("rebeldia", "promocion", "civil", ["rebeldia", "preclusion", "emplazamiento"]),
        ("prevencion", "cumplimiento de prevencion", "civil", ["prevencion", "cumplimiento", "subsanacion"]),
        ("oficio", "oficio", "administrativo", ["oficio", "autoridad", "requerimiento"]),
        ("audio", "promocion", "penal", ["audio", "video", "audiencia"]),
        ("reconsideracion", "recurso", "administrativo", ["reconsideracion", "recurso", "revocacion"]),
        ("recosideracion", "recurso", "administrativo", ["reconsideracion", "recurso", "revocacion"]),
    ]
    for needle, kind, matter, keywords in rules:
        if needle in haystack:
            return kind, matter, keywords
    return "promocion", "general", ["escrito", "promocion", "solicitud"]


def build_manifest_entry(
    source: Path,
    index: int,
    sanitized_docx: Path,
    text_path: Path,
    counts: Counter,
    name_patterns: list[re.Pattern[str]],
) -> dict:
    text = extract_docx_text(sanitized_docx)
    text = repair_mojibake(text)
    placeholder_text, placeholder_counts = redact_text(text, name_patterns, placeholder=True)
    counts.update({f"placeholder_{key}": value for key, value in placeholder_counts.items()})
    kind, matter, keywords = infer_kind_and_matter(source.stem, placeholder_text)
    title = f"Ejemplo {index:02d} - {kind.title()}"
    text_path.write_text(placeholder_text, encoding="utf-8")
    return {
        "id": f"escrito-{index:02d}-{slugify(kind)}",
        "title": title,
        "kind": kind,
        "matter": matter,
        "keywords": sorted(set(keywords + [kind, matter])),
        "sanitizedDocx": f"assets/drafting-examples/docx/{sanitized_docx.name}",
        "sanitizedText": f"assets/drafting-examples/text/{text_path.name}",
        "sourceSha256": sha256_file(source),
        "redactionCounts": dict(sorted(counts.items())),
        "sampleText": placeholder_text[:9000],
    }


def write_typescript_data(entries: list[dict], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(entries, ensure_ascii=False, indent=2)
    output_path.write_text(
        "export interface DraftingExample {\n"
        "  id: string;\n"
        "  title: string;\n"
        "  kind: string;\n"
        "  matter: string;\n"
        "  keywords: string[];\n"
        "  sanitizedDocx: string;\n"
        "  sanitizedText: string;\n"
        "  sourceSha256: string;\n"
        "  redactionCounts: Record<string, number>;\n"
        "  sampleText: string;\n"
        "}\n\n"
        f"export const draftingExamples: DraftingExample[] = {payload};\n",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Sanitize legal drafting examples for Juris/Escriba.")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", default=Path("assets/drafting-examples"), type=Path)
    parser.add_argument("--data", default=Path("src/data/draftingExamples.ts"), type=Path)
    args = parser.parse_args()

    source_dir = args.source.resolve()
    output_dir = args.output.resolve()
    docx_dir = output_dir / "docx"
    text_dir = output_dir / "text"
    report_dir = output_dir / "reports"

    if not source_dir.exists():
        raise SystemExit(f"No existe la carpeta fuente: {source_dir}")

    shutil.rmtree(output_dir, ignore_errors=True)
    docx_dir.mkdir(parents=True, exist_ok=True)
    text_dir.mkdir(parents=True, exist_ok=True)
    report_dir.mkdir(parents=True, exist_ok=True)

    files = sorted(path for path in source_dir.rglob("*") if path.is_file())
    docx_files = [path for path in files if path.suffix.lower() == ".docx"]
    skipped = [
        {
            "fileName": f"archivo-{index:02d}{path.suffix.lower()}",
            "reason": "Formato .doc antiguo pendiente de conversion a .docx" if path.suffix.lower() == ".doc" else "Formato no soportado",
        }
        for index, path in enumerate(files, start=1)
        if path.suffix.lower() != ".docx"
    ]
    filename_terms = infer_filename_terms(files)
    name_patterns = make_name_patterns(filename_terms)

    entries = []
    for index, source in enumerate(docx_files, start=1):
        kind, matter, _ = infer_kind_and_matter(source.stem, "")
        slug = f"escrito-{index:02d}-{slugify(kind)}-{slugify(matter)}"
        sanitized_docx = docx_dir / f"{slug}.docx"
        text_path = text_dir / f"{slug}.txt"
        counts = sanitize_docx(source, sanitized_docx, name_patterns)
        entries.append(build_manifest_entry(source, index, sanitized_docx, text_path, counts, name_patterns))

    manifest = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceDirectory": "Base de Datos JM",
        "processedCount": len(entries),
        "skipped": skipped,
        "entries": entries,
    }
    (report_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    write_typescript_data(entries, args.data.resolve())

    print(json.dumps({
        "processed": len(entries),
        "skipped": skipped,
        "output": str(output_dir),
        "data": str(args.data.resolve()),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
